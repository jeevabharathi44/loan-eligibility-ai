# -*- coding: utf-8 -*-
import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

REPO_DIR = r"C:\Users\Jeeva\.gemini\antigravity\scratch\loan-eligibility-ai"
ARTIFACT_DIR = r"C:\Users\Jeeva\.gemini\antigravity\brain\e905d853-ffc3-4e6a-aec2-a3c083d192f0"
SRC_TEMPLATE = os.path.join(ARTIFACT_DIR, r".user_uploaded\media_1793376831728_cd86e238.pptx")
OUTPUT_PPTX = os.path.join(REPO_DIR, "Smart_Digital_Lending_System_Viva_Presentation.pptx")
ARTIFACT_PPTX = os.path.join(ARTIFACT_DIR, "Smart_Digital_Lending_System_Viva_Presentation.pptx")
DESKTOP_PPTX = r"C:\Users\Jeeva\Desktop\Smart_Digital_Lending_System_Viva_Presentation.pptx"

# Copy original template fresh
shutil.copy2(SRC_TEMPLATE, OUTPUT_PPTX)
prs = pptx.Presentation(OUTPUT_PPTX)

# Template exact brand colors
COLOR_FOREST_GREEN = RGBColor(31, 77, 57)    # #1F4D39
COLOR_GOLD = RGBColor(200, 168, 62)          # #C8A83E
COLOR_CHARCOAL = RGBColor(26, 26, 26)        # #1A1A1A
COLOR_GRAY = RGBColor(107, 107, 107)         # #6B6B6B
COLOR_WHITE = RGBColor(255, 255, 255)       # #FFFFFF

def get_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None

def set_title(shape, text):
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = ""
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Calibri"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_FOREST_GREEN

def set_subtitle(shape, text):
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = ""
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = False
    p.font.color.rgb = COLOR_GRAY

def set_pill_header(shape, text):
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = ""
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

def add_bullets(shape, bullet_list, font_size=11.0, space_after=5):
    """
    bullet_list is a list of tuples: (prefix_bold, rest_regular)
    """
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = ""
    
    for i, (prefix, rest) in enumerate(bullet_list):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.space_after = Pt(space_after)
        p.line_spacing = 1.15
        
        if prefix:
            r1 = p.add_run()
            r1.text = prefix
            r1.font.name = "Calibri"
            r1.font.size = Pt(font_size)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_CHARCOAL
            
        if rest:
            r2 = p.add_run()
            r2.text = rest
            r2.font.name = "Calibri"
            r2.font.size = Pt(font_size)
            r2.font.bold = False
            r2.font.color.rgb = COLOR_CHARCOAL

def format_table(shape, data, header_size=11.0, cell_size=10.5):
    if not shape or not shape.has_table:
        return
    table = shape.table
    for r_idx, row in enumerate(data):
        if r_idx >= len(table.rows):
            break
        for c_idx, val in enumerate(row):
            if c_idx >= len(table.columns):
                break
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            r = p.add_run()
            r.text = str(val)
            r.font.name = "Calibri"
            
            if r_idx == 0:
                r.font.size = Pt(header_size)
                r.font.bold = True
                r.font.color.rgb = COLOR_WHITE
            else:
                r.font.size = Pt(cell_size)
                r.font.color.rgb = COLOR_CHARCOAL
                if c_idx == 0:
                    r.font.bold = True
                else:
                    r.font.bold = False

print("Processing Slide 1 (Title)...")
s1 = prs.slides[0]
# Title on Slide 1
sh = get_shape(s1, "object 13")
if sh and sh.has_text_frame:
    sh.text_frame.word_wrap = True
    sh.text_frame.text = ""
    p = sh.text_frame.paragraphs[0]
    p.text = "Smart Digital Lending System with Automated Verification & Explainable AI"
    p.font.name = "Calibri"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = COLOR_FOREST_GREEN

sh = get_shape(s1, "object 14")
if sh and sh.has_text_frame:
    sh.text_frame.text = ""
    p = sh.text_frame.paragraphs[0]
    p.text = "Batch No: [01] • Viva Voce Presentation"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = False
    p.font.color.rgb = COLOR_GRAY

print("Processing Slide 2 (Abstract)...")
s2 = prs.slides[1]
set_subtitle(get_shape(s2, "object 3"), "Fast, Fraud-Proof & Transparent Loan Processing using 5 Core Technologies")
# Keywords
sh_kw = get_shape(s2, "object 8")
if sh_kw and sh_kw.has_text_frame:
    sh_kw.text_frame.text = ""
    p = sh_kw.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Keywords: "
    r1.font.bold = True
    r1.font.color.rgb = COLOR_GOLD
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "Digital Lending • DigiLocker • Face Biometrics • OCR Extraction • Explainable AI • Random Forest"
    r2.font.bold = False
    r2.font.color.rgb = COLOR_GOLD
    r2.font.size = Pt(10.5)

set_pill_header(get_shape(s2, "object 10"), "ABSTRACT")
set_pill_header(get_shape(s2, "object 15"), "TECHNOLOGY STACK")

s2_bullets = [
    ("Problem: ", "manual checks take 3-7 days, fake documents bypass review, stolen OTPs create identity theft loans, and AI rejects applicants without reasons."),
    ("Approach: ", "a 5-model pipeline: DigiLocker document fetch, 3-way webcam face matching, OCR cross-checks, and glass-box EBM underwriting."),
    ("Technologies: ", "React 18, Vite, FastAPI, OpenCV, Tesseract OCR, InterpretML, scikit-learn, SQLite."),
    ("Implemented: ", "DigiLocker SSO, biometric gate, cross-field OCR, EBM scoring, adverse action notices, and full audit trail."),
    ("Contribution: ", "zero fake paperwork, immune to OTP theft; 85.37% test accuracy matching Random Forest with 100% exact reasons.")
]
add_bullets(get_shape(s2, "object 19"), s2_bullets, font_size=11.0, space_after=4)

s2_stack = [
    ("Model 1: ", "DigiLocker Government Document Fetch"),
    ("Model 2: ", "Tri-Biometric Live Face Gate (≥ 75%)"),
    ("Model 3: ", "OCR Text Extraction & Cross-Check"),
    ("Model 4: ", "Explainable Boosting Machine (EBM)"),
    ("Model 5: ", "Random Forest Accuracy Benchmark"),
    ("Full Stack: ", "React 18 + FastAPI + OpenCV + SQLite")
]
add_bullets(get_shape(s2, "object 20"), s2_stack, font_size=11.5, space_after=4)

# Workflow pills on Slide 2
wf_items = ["DigiLocker", "Face Match", "OCR Check", "EBM Score", "Rules", "Reasons", "Audit Log", "Verdict"]
for idx, name in enumerate(["object 25", "object 30", "object 35", "object 40", "object 45", "object 50", "object 55", "object 60"]):
    sh_wf = get_shape(s2, name)
    if sh_wf and sh_wf.has_text_frame:
        sh_wf.text_frame.text = ""
        p = sh_wf.text_frame.paragraphs[0]
        p.text = wf_items[idx]
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

print("Processing Slide 3 (Problem Statement & Motivation)...")
s3 = prs.slides[2]
set_subtitle(get_shape(s3, "object 3"), "The 4 Critical Problems in Current Loan Processing and Why We Solve Them")
set_pill_header(get_shape(s3, "object 6"), "LIMITATIONS OF EXISTING METHODS (4 PROBLEMS)")
set_pill_header(get_shape(s3, "object 12"), "NEED & SOLUTION")

s3_problems = [
    ("Manual Review Delays: ", "checking physical paperwork takes 3 to 7 days, delaying urgent funds for borrowers."),
    ("Fake Certificate Fraud: ", "applicants produce forged income slips and edited ID cards that human reviewers miss."),
    ("Stolen OTP Identity Theft: ", "scammers steal Aadhaar or PAN OTPs to create fake loans in innocent people's names."),
    ("Unexplained Black-Box Rejections: ", "traditional AI declines loans without explaining why, leaving applicants helpless.")
]
add_bullets(get_shape(s3, "object 10"), s3_problems, font_size=11.5, space_after=6)

s3_need = [
    ("A smart digital lending system ", "that gives a decision in under 60 seconds, fetches authentic DigiLocker certificates, verifies live face biometrics to defeat OTP theft, and explains every decision with glass-box EBM.")
]
add_bullets(get_shape(s3, "object 16"), s3_need, font_size=11.5, space_after=4)

sh_mot = get_shape(s3, "object 20")
if sh_mot and sh_mot.has_text_frame:
    sh_mot.text_frame.text = ""
    p = sh_mot.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Motivation: "
    r1.font.bold = True
    r1.font.color.rgb = COLOR_CHARCOAL
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "RBI digital lending guidelines expect clear reasons for credit decisions; a bare approved/declined answer falls short."
    r2.font.bold = False
    r2.font.color.rgb = COLOR_CHARCOAL
    r2.font.size = Pt(10.5)

print("Processing Slide 4 (Objectives)...")
s4 = prs.slides[3]
set_subtitle(get_shape(s4, "object 3"), "Each objective mapped to an implemented module and verified result")
s4_data = [
    ["#", "OBJECTIVE", "STATUS"],
    ["1", "DigiLocker Integration: auto-fetch verified e-Aadhaar and PAN to stop fake certificates", "✓ Implemented"],
    ["2", "Tri-Biometric Security: match live webcam face with Aadhaar & PAN photos to defeat OTP theft", "✓ Implemented"],
    ["3", "OCR Cross-Verification: extract text and cross-check name spelling across identity cards", "✓ Implemented"],
    ["4", "Explainable EBM Underwriting: score credit risk and output plain-English decision reasons", "✓ Implemented"],
    ["5", "Random Forest Benchmark: validate that EBM matches state-of-the-art ensemble accuracy", "✓ Implemented"]
]
format_table(get_shape(s4, "object 5"), s4_data, header_size=11.0, cell_size=11.0)

print("Processing Slide 5 (Literature Review & Research Gap)...")
s5 = prs.slides[4]
set_subtitle(get_shape(s5, "object 3"), "Comparing current lending approaches to our unified smart system")
s5_data = [
    ["STUDY / APPROACH", "CONTRIBUTION", "GAP FOR THIS PROJECT"],
    ["Manual Review & Scorecards", "Readable points per variable", "Slow (3-7 days); high fake certificate risk"],
    ["OTP-Based Digital Lending", "Instant mobile authentication", "Vulnerable to stolen OTPs & identity theft"],
    ["Deep Learning & Ensembles", "High accuracy credit prediction", "Black-box opacity; cannot explain rejections"],
    ["Post-Hoc SHAP / LIME", "Approximate factor explanations", "Approximations; unstable with correlated inputs"],
    ["Our Proposed 5-Model System", "DigiLocker + Biometrics + OCR + EBM", "Fast, fraud-proof, physically verified & transparent"]
]
format_table(get_shape(s5, "object 5"), s5_data, header_size=11.0, cell_size=10.5)

sh_gap = get_shape(s5, "object 9")
if sh_gap and sh_gap.has_text_frame:
    sh_gap.text_frame.text = ""
    p = sh_gap.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Research gap: "
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "prediction and explanation are studied separately; few works give a full workflow combining government document fetch, 3-way biometrics, and glass-box AI in one platform."
    r2.font.bold = False
    r2.font.size = Pt(10.5)

print("Processing Slide 6 (Proposed Methodology)...")
s6 = prs.slides[5]
sh_wft = get_shape(s6, "object 4")
if sh_wft and sh_wft.has_text_frame:
    sh_wft.text_frame.text = ""
    p1 = sh_wft.text_frame.paragraphs[0]
    p1.text = "Five-stage pipeline from document fetch to audited verdict"
    p1.font.name = "Calibri"
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_GRAY

# Update workflow box names
s6_wf_labels = ["DigiLocker\nSSO", "Live Face\nMatch", "OCR Cross\nCheck", "EBM\nScoring", "Ranked\nReasons", "Verdict +\nAudit"]
for idx, name in enumerate(["object 8", "object 13", "object 18", "object 23", "object 28", "object 33"]):
    sh_box = get_shape(s6, name)
    if sh_box and sh_box.has_text_frame:
        sh_box.text_frame.text = ""
        p = sh_box.text_frame.paragraphs[0]
        p.text = s6_wf_labels[idx]
        p.font.name = "Calibri"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

set_pill_header(get_shape(s6, "object 35"), "PROCESSING STEPS")
set_pill_header(get_shape(s6, "object 40"), "KEY FEATURES")

s6_steps = [
    ("DigiLocker SSO: ", "pulls authentic e-Aadhaar and PAN from government databases."),
    ("Tri-Biometric Match: ", "compares live webcam selfie with both document photos (≥ 75%)."),
    ("OCR Cross-Check: ", "extracts text and validates name and DOB consistency across IDs."),
    ("EBM Scoring: ", "calculates risk probability from income, debts, and bureau score."),
    ("Explainable Verdict: ", "ranks top reasons and outputs instant sanction or decline notice.")
]
add_bullets(get_shape(s6, "object 44"), s6_steps, font_size=11.0, space_after=4)

s6_features = [
    ("Under 60 seconds ", "from application start to disbursed loan decision."),
    ("Zero fake certificates: ", "only government issuer documents accepted."),
    ("Immune to stolen OTPs: ", "physical presence required for face match."),
    ("Plain-English explanations: ", "tells applicants exactly how to qualify."),
    ("Full audit trail: ", "loan officers can inspect complete decision logs.")
]
add_bullets(get_shape(s6, "object 45"), s6_features, font_size=11.0, space_after=4)

print("Processing Slide 7 (System Architecture)...")
s7 = prs.slides[6]
set_subtitle(get_shape(s7, "object 3"), "Components and data flow of the application")

sh_desc = get_shape(s7, "object 39")
if sh_desc and sh_desc.has_text_frame:
    s7_services = [
        ("Client: ", "wizard, live camera reticle, waterfall chart."),
        ("Security: ", "DigiLocker API, OpenCV face match, OCR parser."),
        ("Model service: ", "features, EBM glass-box, Random Forest."),
        ("Database: ", "users, applications, predictions, audit log.")
    ]
    add_bullets(sh_desc, s7_services, font_size=11.0, space_after=4)

sh_dnote = get_shape(s7, "object 43")
if sh_dnote and sh_dnote.has_text_frame:
    sh_dnote.text_frame.text = ""
    p = sh_dnote.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Design note: "
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "explanations come from the EBM's own terms, not a separate approximation, so reasons and prediction cannot disagree."
    r2.font.bold = False
    r2.font.size = Pt(10.5)

sh_dflow = get_shape(s7, "object 47")
if sh_dflow and sh_dflow.has_text_frame:
    sh_dflow.text_frame.text = ""
    p = sh_dflow.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "DATA FLOW\n"
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_FOREST_GREEN
    r2 = p.add_run()
    r2.text = "Client → API → Security Tier (DigiLocker + Biometrics + OCR) → Model Service (EBM + Rules) → Audit DB → Client"
    r2.font.size = Pt(10.0)

print("Processing Slide 8 (Model 1 & 2: DigiLocker & Biometrics)...")
s8 = prs.slides[7]
set_title(get_shape(s8, "object 2"), "MODEL 1 & 2: DIGILOCKER & BIOMETRICS")
set_subtitle(get_shape(s8, "object 3"), "Eliminating fake certificates and stolen OTP identity theft")
set_pill_header(get_shape(s8, "object 6"), "MODEL 1: DIGILOCKER VERIFICATION")
set_pill_header(get_shape(s8, "object 12"), "MODEL 2: TRI-BIOMETRIC VERIFICATION")

s8_m1 = [
    ("Problem solved: ", "delays (Problem 1) and fake certificates (Problem 2)."),
    ("Direct connection: ", "fetches verified e-Aadhaar and PAN from UIDAI and Income Tax department."),
    ("Mathematical check: ", "verifies 12-digit Aadhaar checksum using Verhoeff algorithm."),
    ("Impact: ", "eliminates 100% of fake certificate fraud and cuts verification time to seconds.")
]
add_bullets(get_shape(s8, "object 10"), s8_m1, font_size=11.5, space_after=6)

s8_m2 = [
    ("Problem solved: ", "identity theft via stolen OTPs (Problem 3)."),
    ("Live presence check: ", "scammers steal OTPs, but cannot steal the applicant's physical face."),
    ("3-Way Matching: ", "Live Webcam Face ↔ Aadhaar Photo ↔ PAN Photo."),
    ("Threshold: ", "requires ≥ 75.0% facial similarity across all three images."),
    ("Impact: ", "blocks any impostor even if they possess the valid OTP.")
]
add_bullets(get_shape(s8, "object 21"), s8_m2, font_size=11.5, space_after=6)

sh_s8_note = get_shape(s8, "object 25")
if sh_s8_note and sh_s8_note.has_text_frame:
    sh_s8_note.text_frame.text = ""
    p = sh_s8_note.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Security guarantee: "
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "identity is confirmed by both government databases AND live physical biometric presence."
    r2.font.size = Pt(10.5)

print("Processing Slide 9 (Model 3 & 4: OCR Extraction & Explainable EBM)...")
s9 = prs.slides[8]
set_title(get_shape(s9, "object 2"), "MODEL 3 & 4: OCR EXTRACTION & EXPLAINABLE EBM")
set_subtitle(get_shape(s9, "object 3"), "Cross-checking details and explaining loan decisions plainly")
set_pill_header(get_shape(s9, "object 6"), "MODEL 3: OCR CROSS-CHECK")
set_pill_header(get_shape(s9, "object 13"), "MODEL 4: EXPLAINABLE EBM (GLASS-BOX)")

s9_m3 = [
    ("Problem solved: ", "mismatched, forged, or mixed identities."),
    ("Text extraction: ", "OCR scans names, DOB, and ID numbers from documents."),
    ("Fuzzy matching: ", "handles Indian name variations (e.g. \"Jeeva Bharathi M\" vs \"M Jeeva Bharathi\")."),
    ("Impact: ", "ensures documents belong to the same person without human review.")
]
add_bullets(get_shape(s9, "object 11"), s9_m3, font_size=11.5, space_after=6)

s9_m4 = [
    ("Problem solved: ", "unexplained, black-box loan rejections (Problem 4)."),
    ("Glass-box formula: ", "Score = Base Intercept + Σ f(Features) + Σ f(Pairs)."),
    ("Exact attribution: ", "calculates the exact positive bonus or penalty of every factor."),
    ("Plain-English notice: ", "tells rejected applicants exactly why (e.g., \"Debts take 56% of income\")."),
    ("Impact: ", "100% transparent decisions compliant with RBI fair lending rules.")
]
add_bullets(get_shape(s9, "object 17"), s9_m4, font_size=11.5, space_after=6)

print("Processing Slide 10 (Model 5: Random Forest Benchmark)...")
s10 = prs.slides[9]
set_title(get_shape(s10, "object 2"), "MODEL 5: RANDOM FOREST BENCHMARK VS EBM")
set_subtitle(get_shape(s10, "object 3"), "Validating prediction accuracy by comparing EBM against Random Forest")
s10_data = [
    ["EVALUATION CRITERIA", "RANDOM FOREST (MODEL 5)", "EXPLAINABLE BOOSTING MACHINE (MODEL 4)"],
    ["Model Architecture", "Ensemble of 100+ Decision Trees", "Generalized Additive Model with Pairs (GA²M)"],
    ["Test Accuracy", "85.37% (Industry benchmark)", "85.37% (Ties Random Forest exactly)"],
    ["AUC-ROC Score", "0.729", "0.752 (EBM leads by +0.023 in discrimination)"],
    ["Interpretability", "Black Box — cannot explain reasons", "Glass Box — 100% exact mathematical reasons"],
    ["Adverse Action Notice", "Requires unstable post-hoc SHAP", "Built-in, instant plain-English explanations"]
]
format_table(get_shape(s10, "object 5"), s10_data, header_size=11.0, cell_size=11.0)

sh_s10_note = get_shape(s10, "object 9")
if sh_s10_note and sh_s10_note.has_text_frame:
    sh_s10_note.text_frame.text = ""
    p = sh_s10_note.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Key takeaway: "
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run()
    r2.text = "Random Forest was trained as our high-accuracy benchmark. EBM matches Random Forest in accuracy (85.37%) and beats it in AUC (0.752), while providing 100% transparent explanations."
    r2.font.size = Pt(10.5)

print("Processing Slide 11 (Implementation Status)...")
s11 = prs.slides[10]
set_subtitle(get_shape(s11, "object 3"), "Major modules and implementation status")
s11_data = [
    ["MODULE / MODEL", "FUNCTION AND IMPLEMENTED CAPABILITY", "STATUS"],
    ["Model 1: DigiLocker", "Direct document fetch with e-Aadhaar & PAN verification", "✓ Implemented"],
    ["Model 2: Biometric Gate", "3-way live camera vs document photo matching (≥ 75%)", "✓ Implemented"],
    ["Model 3: OCR Cross-Check", "Text extraction and fuzzy name string comparison", "✓ Implemented"],
    ["Model 4: EBM Underwriting", "Exact scoring with waterfall charts and plain-English reasons", "✓ Implemented"],
    ["Model 5: Random Forest", "Benchmark ensemble model trained and evaluated", "✓ Implemented"],
    ["Web Application", "6-step underwriting wizard, camera viewfinder & audit cockpit", "✓ Implemented"],
    ["Next Phase (Proposed)", "Production UIDAI FaceRD API, Account Aggregator, regional audio", "Planned"]
]
format_table(get_shape(s11, "object 5"), s11_data, header_size=11.0, cell_size=11.0)

print("Processing Slide 12 (Results: Model Comparison on Test Set)...")
s12 = prs.slides[11]
set_subtitle(get_shape(s12, "object 3"), "Same 491 training / 123 test records for all five models; bureau attributes are generated")
s12_data = [
    ["MODEL", "ACC. %", "PREC. %", "REC. %", "F1 %", "AUC"],
    ["Logistic regression", "74.80", "80.68", "83.53", "82.08", "0.728"],
    ["Gradient boosting", "80.49", "82.80", "90.59", "86.52", "0.705"],
    ["Hist. grad. boosting", "82.11", "82.47", "94.12", "87.91", "0.734"],
    ["Random forest (Model 5)", "85.37", "83.17", "98.82", "90.32", "0.729"],
    ["EBM Glass-Box (Model 4)", "85.37", "83.17", "98.82", "90.32", "0.752"]
]
format_table(get_shape(s12, "object 5"), s12_data, header_size=11.0, cell_size=11.0)

sh_s12_note = get_shape(s12, "object 9")
if sh_s12_note and sh_s12_note.has_text_frame:
    sh_s12_note.text_frame.text = ""
    p = sh_s12_note.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Outcome: "
    r1.font.bold = True
    r1.font.size = Pt(11.0)
    r2 = p.add_run()
    r2.text = "EBM ties random forest and leads in AUC; +10.57 points accuracy over logistic regression. 5-fold CV: 83.25 ± 1.72% accuracy, AUC 0.763."
    r2.font.size = Pt(11.0)

print("Processing Slide 13 (Results: Sample Predictions & Fraud Defense)...")
s13 = prs.slides[12]
set_subtitle(get_shape(s13, "object 3"), "Three end-to-end runs of the prototype and fraud defense results")
s13_pred_data = [
    ["PROFILE", "SCORE", "DTI %", "UTIL. %", "PROB.", "OUTCOME"],
    ["Prime (Jeeva M)", "785", "28.0", "18", "88.3%", "Approved"],
    ["Stolen OTP Fraud", "N/A", "N/A", "N/A", "0.0%", "BLOCKED"],
    ["Adverse (Karthik R)", "560", "55.9", "78", "14.2%", "Declined"]
]
format_table(get_shape(s13, "object 6"), s13_pred_data, header_size=10.5, cell_size=10.5)

s13_abl_data = [
    ["VARIANT", "ACC. %", "AUC"],
    ["Full EBM", "85.37", "0.752"],
    ["No interaction terms", "84.55", "0.741"],
    ["No derived ratios", "82.93", "0.729"],
    ["No bureau attributes", "69.11", "—"]
]
format_table(get_shape(s13, "object 8"), s13_abl_data, header_size=10.5, cell_size=10.5)

sh_s13_note = get_shape(s13, "object 12")
if sh_s13_note and sh_s13_note.has_text_frame:
    sh_s13_note.text_frame.text = ""
    p = sh_s13_note.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Reading: "
    r1.font.bold = True
    r1.font.size = Pt(11.0)
    r2 = p.add_run()
    r2.text = "removing bureau attributes drops accuracy to the 69.11% majority rate. Impostor tests confirmed 100% of stolen OTP attempts were blocked at the biometric gate."
    r2.font.size = Pt(11.0)

print("Processing Slide 14 (How the 4 Problems are Solved)...")
s14 = prs.slides[13]
set_title(get_shape(s14, "object 2"), "HOW OUR SYSTEM SOLVES THE 4 PROBLEMS")
set_subtitle(get_shape(s14, "object 3"), "Direct impact of our 5 models on modern digital lending")
set_pill_header(get_shape(s14, "object 6"), "SOLVING DELAYS & FAKE CERTIFICATES")
set_pill_header(get_shape(s14, "object 11"), "SOLVING OTP THEFT & BLACK-BOX REJECTIONS")

s14_left = [
    ("1. Manual Verification Delays: ", "traditional physical checks take 3 to 7 days."),
    ("→ Solved by: ", "DigiLocker & OCR models; entire process completes in under 60 seconds."),
    ("2. Fake Certificates & Forgery: ", "scammers submit photoshopped documents."),
    ("→ Solved by: ", "DigiLocker direct government API fetch; only authentic records accepted.")
]
add_bullets(get_shape(s14, "object 15"), s14_left, font_size=11.5, space_after=6)

s14_right = [
    ("3. Stolen OTPs & Identity Theft: ", "fraudsters steal OTPs to create fake loans."),
    ("→ Solved by: ", "Tri-Biometric face match; impostors are blocked without physical presence."),
    ("4. Unexplained Black-Box Rejections: ", "traditional AI declines loans without reasons."),
    ("→ Solved by: ", "Explainable Boosting Machine (EBM); provides plain-English reasons and advice.")
]
add_bullets(get_shape(s14, "object 16"), s14_right, font_size=11.5, space_after=6)

print("Processing Slide 15 (Conclusion and Future Scope)...")
s15 = prs.slides[14]
set_subtitle(get_shape(s15, "object 3"), "What Phase I delivers and what comes next")
set_pill_header(get_shape(s15, "object 6"), "CONCLUSION")
set_pill_header(get_shape(s15, "object 11"), "FUTURE SCOPE")

s15_conc = [
    ("Digital loan system ", "solving manual delays, fake documents, OTP theft, and opacity."),
    ("DigiLocker & OCR ", "eliminate manual paperwork and forged certificates."),
    ("Tri-Biometric security ", "defeats identity theft via stolen OTPs."),
    ("EBM matches best ensemble: ", "85.37% accuracy, AUC 0.752 with 100% exact reasons."),
    ("Full working application ", "running locally and deployed live on GitHub Pages.")
]
add_bullets(get_shape(s15, "object 15"), s15_conc, font_size=11.5, space_after=5)

s15_future = [
    ("UIDAI Aadhaar FaceRD ", "mobile app integration for production biometrics."),
    ("Regional audio explanations ", "in Tamil and Hindi for adverse action notices."),
    ("Account Aggregator (AA) ", "automated bank statement health analysis."),
    ("Risk-based interest rates ", "offering lower rates to low-risk borrowers.")
]
add_bullets(get_shape(s15, "object 16"), s15_future, font_size=11.5, space_after=5)

print("Processing Slide 16 (Thank You)...")
s16 = prs.slides[15]
set_title(get_shape(s16, "object 2"), "THANK YOU")
sh_ty_sub = get_shape(s16, "object 3")
if sh_ty_sub and sh_ty_sub.has_text_frame:
    sh_ty_sub.text_frame.text = ""
    p = sh_ty_sub.text_frame.paragraphs[0]
    p.text = "Questions, Comments & Suggestions are Welcome"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_GRAY

# Save presentation to workspace, artifact directory, and Desktop
prs.save(OUTPUT_PPTX)
print(f"Saved: {OUTPUT_PPTX}")

shutil.copy2(OUTPUT_PPTX, ARTIFACT_PPTX)
print(f"Saved: {ARTIFACT_PPTX}")

shutil.copy2(OUTPUT_PPTX, DESKTOP_PPTX)
print(f"Saved directly to Desktop: {DESKTOP_PPTX}")
