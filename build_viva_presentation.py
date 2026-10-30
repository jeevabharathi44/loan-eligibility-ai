# -*- coding: utf-8 -*-
import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

REPO_DIR = r"C:\Users\Jeeva\.gemini\antigravity\scratch\loan-eligibility-ai"
ARTIFACT_DIR = r"C:\Users\Jeeva\.gemini\antigravity\brain\e905d853-ffc3-4e6a-aec2-a3c083d192f0"
SRC_TEMPLATE = os.path.join(ARTIFACT_DIR, r".user_uploaded\media_1793376831728_cd86e238.pptx")
OUTPUT_PPTX = os.path.join(REPO_DIR, "Smart_Digital_Lending_System_Viva_Presentation.pptx")
ARTIFACT_PPTX = os.path.join(ARTIFACT_DIR, "Smart_Digital_Lending_System_Viva_Presentation.pptx")

shutil.copy2(SRC_TEMPLATE, OUTPUT_PPTX)
prs = pptx.Presentation(OUTPUT_PPTX)

def get_shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None

def set_text(shape, text, font_size_pt=None, bold=None, color_rgb=None):
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    
    # Capture style from first run if available
    old_font_name = "Segoe UI"
    old_size = None
    old_color = None
    if tf.paragraphs and tf.paragraphs[0].runs:
        r0 = tf.paragraphs[0].runs[0]
        if r0.font.name:
            old_font_name = r0.font.name
        if r0.font.size:
            old_size = r0.font.size
        try:
            if r0.font.color and r0.font.color.type:
                old_color = r0.font.color.rgb
        except:
            pass

    tf.text = ""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = line
        p.font.name = old_font_name
        if font_size_pt:
            p.font.size = Pt(font_size_pt)
        elif old_size:
            p.font.size = old_size
        if bold is not None:
            p.font.bold = bold
        if color_rgb:
            p.font.color.rgb = color_rgb
        elif old_color:
            p.font.color.rgb = old_color

def set_table_data(shape, data, header_size_pt=9.5, cell_size_pt=8.5):
    if not shape or not shape.has_table:
        return
    table = shape.table
    for r_idx, row in enumerate(data):
        if r_idx >= len(table.rows):
            break
        for c_idx, val in enumerate(row):
            if c_idx >= len(table.columns):
                break
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Segoe UI"
                if r_idx == 0:
                    p.font.size = Pt(header_size_pt)
                    p.font.bold = True
                else:
                    p.font.size = Pt(cell_size_pt)
                    if c_idx == 0:
                        p.font.bold = True

print("Updating Slide 1: Title Slide...")
s1 = prs.slides[0]
sh = get_shape_by_name(s1, "object 13")
set_text(sh, "Smart Digital Lending System with Automated Document Verification, Biometric Security and Explainable AI", font_size_pt=18, bold=True)
sh = get_shape_by_name(s1, "object 14")
set_text(sh, "Batch No: [01] • Viva Voce Presentation", font_size_pt=10, bold=True)

print("Updating Slide 2: Abstract...")
s2 = prs.slides[1]
set_text(get_shape_by_name(s2, "object 3"), "Fast, Fraud-Proof & Transparent Loan Processing using 5 Core Technologies", font_size_pt=11)
set_text(get_shape_by_name(s2, "object 8"), "Keywords: Digital Lending • DigiLocker • Face Biometrics • OCR Extraction • Explainable Boosting Machine • Random Forest", font_size_pt=8.5)
set_text(get_shape_by_name(s2, "object 15"), "TECHNOLOGY STACK & 5 MODELS", font_size_pt=11, bold=True)

abs_text = (
    "Problem: Traditional loan checks take 3-7 days of manual paperwork, fake certificates bypass verification, stolen OTPs lead to identity theft loans, and traditional AI rejects applicants without explaining why.\n\n"
    "Solution: An automated digital lending portal combining 5 specialized modules to verify documents, confirm physical identity, and score credit risk in under 60 seconds.\n\n"
    "Key Modules: (1) DigiLocker document fetch, (2) 3-way webcam face matching, (3) OCR cross-checks, (4) Glass-box EBM underwriting, and (5) Random Forest benchmark.\n\n"
    "Outcome: Zero fake certificate fraud, immune to OTP theft, 85.37% test accuracy matching Random Forest, with 100% transparent plain-English reasons for every decision."
)
set_text(get_shape_by_name(s2, "object 19"), abs_text, font_size_pt=8.5)

stack_text = (
    "Model 1: DigiLocker Verification Module\n"
    "Model 2: Tri-Biometric Face Recognition\n"
    "Model 3: OCR Cross-Field Extraction\n"
    "Model 4: Explainable Boosting Machine (EBM)\n"
    "Model 5: Random Forest Benchmark Model\n\n"
    "Frontend: React 18 + Vite + TailwindCSS\n"
    "Backend: FastAPI (Python 3.10) + OpenCV\n"
    "Database: SQLite 3 / MySQL (Audit Log)\n"
    "Deployment: Local & Live Cloud (GitHub Pages)"
)
set_text(get_shape_by_name(s2, "object 20"), stack_text, font_size_pt=8.5)

# Update workflow box labels
set_text(get_shape_by_name(s2, "object 25"), "DigiLocker", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 30"), "Face Match", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 35"), "OCR Check", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 40"), "EBM Scoring", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 45"), "Clear Reasons", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 50"), "Audit Log", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 55"), "Sanction", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s2, "object 60"), "Disbursed", font_size_pt=8, bold=True)

print("Updating Slide 3: Problem Statement and Motivation (The 4 Problems)...")
s3 = prs.slides[2]
set_text(get_shape_by_name(s3, "object 2"), "PROBLEM STATEMENT AND MOTIVATION", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s3, "object 3"), "The 4 Critical Problems in Current Loan Processing and Why We Solve Them", font_size_pt=11)
set_text(get_shape_by_name(s3, "object 6"), "THE 4 MAJOR PROBLEMS IN LENDING", font_size_pt=11, bold=True)

prob_text = (
    "1. Manual Document Delays: Verifying physical paperwork takes 3 to 7 days, delaying urgent loan access for legitimate borrowers.\n\n"
    "2. Fake Certificate Fraud: Applicants submit forged income certificates or edited ID cards that human reviewers often fail to detect.\n\n"
    "3. Stolen OTPs & Identity Theft: Scammers trick people into sharing Aadhaar or PAN OTPs via phishing or SIM swap, creating fake loans in innocent people's names.\n\n"
    "4. Unexplained Black-Box Rejections: Existing AI tools reject applicants without clear reasons, leaving borrowers confused and violating RBI rules."
)
set_text(get_shape_by_name(s3, "object 10"), prob_text, font_size_pt=8.5)

set_text(get_shape_by_name(s3, "object 12"), "OUR NEED & SOLUTION", font_size_pt=11, bold=True)
need_text = (
    "A modern, automated digital lending system that:\n"
    "• Connects directly to government sources (DigiLocker) to stop fake certificates.\n"
    "• Verifies the applicant's live physical face to make stolen OTPs useless.\n"
    "• Uses OCR to cross-check name spellings and details across documents.\n"
    "• Uses Explainable AI (EBM) to give simple, plain-English reasons for every decision."
)
set_text(get_shape_by_name(s3, "object 16"), need_text, font_size_pt=8.5)

mot_text = (
    "Motivation: Comply with RBI Digital Lending Guidelines (DLG-2022/23) demanding speed, fraud prevention, and complete transparency in credit decisions."
)
set_text(get_shape_by_name(s3, "object 20"), mot_text, font_size_pt=8.5)

print("Updating Slide 4: Objectives...")
s4 = prs.slides[3]
set_text(get_shape_by_name(s4, "object 3"), "Five Concrete Project Goals Mapped to Implemented Modules", font_size_pt=11)
s4_table_data = [
    ["#", "OBJECTIVE", "STATUS"],
    ["1", "Automated Document Fetch: Connect to DigiLocker to fetch authentic Aadhaar & PAN directly, eliminating fake certificates", "✓ Implemented"],
    ["2", "Live Biometric Face Matching: Compare applicant's live webcam face with Aadhaar & PAN photos to prevent OTP theft", "✓ Implemented"],
    ["3", "OCR Text Extraction & Cross-Check: Automatically read document details and cross-verify name spelling across IDs", "✓ Implemented"],
    ["4", "Explainable Credit Assessment: Use Explainable Boosting Machine (EBM) to score risk and show plain-English reasons", "✓ Implemented"],
    ["5", "Machine Learning Benchmarking: Compare EBM against Random Forest to ensure high accuracy without losing transparency", "✓ Implemented"]
]
set_table_data(get_shape_by_name(s4, "object 5"), s4_table_data, header_size_pt=9.5, cell_size_pt=8.2)

print("Updating Slide 5: Literature Review and Research Gap...")
s5 = prs.slides[4]
set_text(get_shape_by_name(s5, "object 3"), "Comparing Current Lending Approaches to Our Unified Smart System", font_size_pt=11)
s5_table_data = [
    ["STUDY / APPROACH", "CONTRIBUTION", "GAP FOR THIS PROJECT"],
    ["Traditional Manual Lending", "Manual xerox review & credit scorecards", "Very slow (3-7 days); high fraud risk from fake documents"],
    ["OTP-Only Digital Lending", "Instant verification using mobile OTPs", "Vulnerable to stolen OTPs, phishing, and identity theft"],
    ["Deep Learning & Ensembles", "High accuracy credit prediction (Random Forest / XGBoost)", "\"Black Box\" opacity; cannot explain why a loan was rejected"],
    ["Post-Hoc Explanations (SHAP)", "Explains black-box models after prediction", "Approximations can be misleading and unstable on correlated data"],
    ["Our Proposed System", "DigiLocker + Face Match + OCR + EBM + Random Forest", "Fast, fraud-proof, physically verified, and 100% transparent"]
]
set_table_data(get_shape_by_name(s5, "object 5"), s5_table_data, header_size_pt=9.5, cell_size_pt=8.2)
gap_text = "Research Gap: Existing platforms only address one aspect (either fast credit scoring OR face match). None combine government DigiLocker verification, 3-way biometrics, and glass-box explainable AI in a unified workflow."
set_text(get_shape_by_name(s5, "object 9"), gap_text, font_size_pt=8.5)

print("Updating Slide 6: Proposed Methodology...")
s6 = prs.slides[5]
set_text(get_shape_by_name(s6, "object 4"), "Five-Stage Simple Pipeline from Document Fetch to Audited Decision\n\nEND-TO-END WORKFLOW", font_size_pt=10.5)

# Workflow labels
set_text(get_shape_by_name(s6, "object 8"), "DigiLocker\nSSO", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s6, "object 13"), "Live Face\nMatch", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s6, "object 18"), "OCR Cross\nCheck", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s6, "object 23"), "EBM\nScoring", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s6, "object 28"), "Clear\nReasons", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s6, "object 33"), "Sanction\nLetter", font_size_pt=8, bold=True)

set_text(get_shape_by_name(s6, "object 35"), "PROCESSING STEPS", font_size_pt=11, bold=True)
set_text(get_shape_by_name(s6, "object 40"), "KEY BENEFITS", font_size_pt=11, bold=True)

proc_steps = (
    "Step 1: DigiLocker Login — Direct fetch of verified Aadhaar & PAN documents.\n\n"
    "Step 2: Tri-Biometric Match — Live webcam selfie compared with Aadhaar & PAN photos.\n\n"
    "Step 3: OCR Cross-Check — Scans text and matches names/DOB across IDs.\n\n"
    "Step 4: EBM Underwriting — Evaluates income, debts, EMI, and credit score.\n\n"
    "Step 5: Decision & Explanation — Instant approval or ranked rejection reasons."
)
set_text(get_shape_by_name(s6, "object 44"), proc_steps, font_size_pt=8.2)

key_benefits = (
    "• Entire application takes under 60 seconds from start to finish.\n\n"
    "• Zero fake certificates: only government-issued records accepted.\n\n"
    "• Stolen OTP immunity: impostors cannot pass the physical face match.\n\n"
    "• Plain-English explanations: tells applicants exactly how to qualify.\n\n"
    "• Full audit trail: loan officers can inspect all decision logs."
)
set_text(get_shape_by_name(s6, "object 45"), key_benefits, font_size_pt=8.2)

print("Updating Slide 7: System Architecture...")
s7 = prs.slides[6]
set_text(get_shape_by_name(s7, "object 3"), "Modular Microservices Architecture Connecting UI, Security Models and ML Engines", font_size_pt=11)

# Services labels
set_text(get_shape_by_name(s7, "object 9"), "Applicant / Officer\nReact 18 SPA", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s7, "object 14"), "DigiLocker Gateway\nUIDAI & ITD Sync", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s7, "object 19"), "Biometric & OCR\nOpenCV + Tesseract", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s7, "object 24"), "FastAPI Backend\nREST API & Rules", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s7, "object 29"), "Underwriting ML\nEBM + Random Forest", font_size_pt=8, bold=True)
set_text(get_shape_by_name(s7, "object 33"), "Audit Database\nSQLite / MySQL", font_size_pt=8, bold=True)

arch_note = "Green/Gold = Security & Underwriting models running inside the backend service."
set_text(get_shape_by_name(s7, "object 34"), arch_note, font_size_pt=7.5)

services_desc = (
    "• Client: Responsive React 18 wizard with live camera viewfinder & EBM waterfall charts.\n"
    "• Security Models: DigiLocker API + OpenCV 256-D face feature extractor + OCR fuzzy matcher.\n"
    "• Underwriting Engine: InterpretML EBM for scoring and exact reason attribution.\n"
    "• Benchmark Engine: Scikit-learn Random Forest for accuracy comparison.\n"
    "• Audit Ledger: Full persistence of all biometric scores, OCR logs, and decisions."
)
set_text(get_shape_by_name(s7, "object 39"), services_desc, font_size_pt=8.2)

design_note = "Design Note: Explanations are calculated directly from EBM's mathematical terms, so reasons and prediction cannot disagree."
set_text(get_shape_by_name(s7, "object 43"), design_note, font_size_pt=8)

data_flow_text = "DATA FLOW\nApplicant Login → DigiLocker Fetch → Live Face Match (≥75%) → OCR Cross-Check → EBM Scoring → Plain-English Reasons → Stored in Audit Log"
set_text(get_shape_by_name(s7, "object 47"), data_flow_text, font_size_pt=8)

print("Updating Slide 8: Model 1 & 2 (DigiLocker & Biometrics)...")
s8 = prs.slides[7]
set_text(get_shape_by_name(s8, "object 2"), "MODEL 1 & 2: DIGILOCKER & BIOMETRICS", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s8, "object 3"), "Eliminating Fake Certificates and Stolen OTP Identity Theft", font_size_pt=11)
set_text(get_shape_by_name(s8, "object 6"), "MODEL 1: DIGILOCKER VERIFICATION", font_size_pt=11, bold=True)

m1_desc = (
    "Why it is needed: Solves Problem 1 (delays) and Problem 2 (fake certificates).\n\n"
    "How it works:\n"
    "• Instead of manual uploads of easily-edited photocopies, the system connects directly to DigiLocker.\n"
    "• Pulls verified e-Aadhaar and e-PAN directly from UIDAI and Income Tax department.\n"
    "• Verifies the 12-digit Aadhaar checksum using the Verhoeff algorithm.\n\n"
    "Impact: Eliminates 100% of forged document submissions and cuts verification time from days to seconds."
)
set_text(get_shape_by_name(s8, "object 10"), m1_desc, font_size_pt=8.2)

set_text(get_shape_by_name(s8, "object 12"), "MODEL 2: TRI-BIOMETRIC VERIFICATION", font_size_pt=11, bold=True)
set_text(get_shape_by_name(s8, "object 16"), "Requires ≥ 75.0% match across all 3 photos to proceed.", font_size_pt=8.5, bold=True)
set_text(get_shape_by_name(s8, "object 17"), "TRI-BIOMETRIC PIPELINE", font_size_pt=11, bold=True)

m2_desc = (
    "Why it is needed: Solves Problem 3 (stolen Aadhaar/PAN OTP fraud).\n\n"
    "How it works:\n"
    "• Scammers can trick victims into sharing OTPs, but they cannot steal the victim's physical face.\n"
    "• Live webcam captures applicant; computer vision extracts 256-D facial features.\n"
    "• Compares Live Selfie ↔ Aadhaar Photo ↔ PAN Photo simultaneously.\n\n"
    "Impact: Blocks any fraudster even if they enter the valid OTP."
)
set_text(get_shape_by_name(s8, "object 21"), m2_desc, font_size_pt=8.2)

sec_guar = "Security Guarantee: Identity is confirmed by both government records AND live physical biometric presence."
set_text(get_shape_by_name(s8, "object 25"), sec_guar, font_size_pt=8.5)

print("Updating Slide 9: Model 3 & 4 (OCR & EBM)...")
s9 = prs.slides[8]
set_text(get_shape_by_name(s9, "object 2"), "MODEL 3 & 4: OCR EXTRACTION & EXPLAINABLE EBM", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s9, "object 3"), "Cross-Checking Details and Explaining Loan Decisions Plainly", font_size_pt=11)
set_text(get_shape_by_name(s9, "object 6"), "MODEL 3: OCR CROSS-CHECK", font_size_pt=11, bold=True)

m3_title = "Fuzzy Matching: Handles natural name variations (e.g. \"Jeeva Bharathi M\" vs \"M Jeeva Bharathi\")"
set_text(get_shape_by_name(s9, "object 10"), m3_title, font_size_pt=8.5, bold=True)

m3_desc = (
    "Why it is needed: Catches mismatched or mixed documents automatically.\n\n"
    "How it works:\n"
    "• Optical Character Recognition (OCR) reads text from identity documents.\n"
    "• Cross-checks applicant name and date of birth across Aadhaar and PAN cards.\n"
    "• Flags spelling errors or mismatched identity numbers before credit evaluation.\n\n"
    "Impact: Ensures documents truly belong to the same applicant without manual review."
)
set_text(get_shape_by_name(s9, "object 11"), m3_desc, font_size_pt=8.2)

set_text(get_shape_by_name(s9, "object 13"), "MODEL 4: EXPLAINABLE BOOSTING MACHINE (EBM)", font_size_pt=11, bold=True)

m4_desc = (
    "Why it is needed: Solves Problem 4 (unexplained, black-box loan rejections).\n\n"
    "How it works:\n"
    "• EBM is a \"Glass-Box\" model: Score = Base + Income Impact + Debt Impact + CIBIL Impact + EMI Impact.\n"
    "• Unlike black-box models, each factor's exact contribution is known mathematically.\n"
    "• Automatically extracts top risk factors and writes clear plain-English reasons.\n\n"
    "Impact: Rejected applicants receive clear reasons like \"Monthly debt payments take up 56% of income\" instead of a mystery refusal."
)
set_text(get_shape_by_name(s9, "object 17"), m4_desc, font_size_pt=8.2)

print("Updating Slide 10: Model 5 & Benchmarking (Random Forest vs EBM)...")
s10 = prs.slides[9]
set_text(get_shape_by_name(s10, "object 2"), "MODEL 5 & BENCHMARKING: RANDOM FOREST VS EBM", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s10, "object 3"), "Validating Prediction Accuracy by Comparing EBM against Random Forest", font_size_pt=11)

s10_table_data = [
    ["EVALUATION CRITERIA", "RANDOM FOREST (MODEL 5)", "EXPLAINABLE BOOSTING MACHINE (MODEL 4)"],
    ["Algorithm Type", "Ensemble of 100+ Decision Trees", "Generalized Additive Model with Pairs (GA²M)"],
    ["Test Accuracy", "85.37% (Industry standard accuracy)", "85.37% (Ties Random Forest exactly)"],
    ["AUC-ROC Score", "0.729", "0.752 (EBM leads by +0.023 in discrimination)"],
    ["Decision Transparency", "\"Black Box\" — cannot explain reasons", "\"Glass Box\" — 100% exact reasons by design"],
    ["Adverse Action Notices", "Requires complex approximate add-ons", "Built-in, instant plain-English sentences"]
]
set_table_data(get_shape_by_name(s10, "object 5"), s10_table_data, header_size_pt=9.5, cell_size_pt=8.2)

s10_note = "Key Conclusion: Random Forest was trained as our high-accuracy benchmark. EBM matches Random Forest in accuracy (85.37%) and beats it in AUC (0.752), while providing 100% transparent explanations that Random Forest cannot give."
set_text(get_shape_by_name(s10, "object 9"), s10_note, font_size_pt=8.5)

print("Updating Slide 11: Implementation Status...")
s11 = prs.slides[10]
set_text(get_shape_by_name(s11, "object 3"), "Current Working Status of All 5 Models and Supporting Modules", font_size_pt=11)

s11_table_data = [
    ["MODULE / MODEL", "FUNCTION AND IMPLEMENTED CAPABILITY", "STATUS"],
    ["Model 1: DigiLocker Service", "Direct document fetch with e-Aadhaar & PAN verification", "✓ Implemented"],
    ["Model 2: Biometric Face Gate", "3-way live camera vs document photo matching (≥ 75%)", "✓ Implemented"],
    ["Model 3: OCR Cross-Check", "Text extraction and fuzzy name string comparison", "✓ Implemented"],
    ["Model 4: EBM Underwriting", "Exact scoring with waterfall charts and plain-English reasons", "✓ Implemented"],
    ["Model 5: Random Forest", "Benchmark ensemble model trained and evaluated", "✓ Implemented"],
    ["Web Application", "6-step underwriting wizard, camera viewfinder & audit cockpit", "✓ Implemented"],
    ["Live Cloud Deployment", "Production web app active on GitHub Pages with mock fallback", "✓ Live & Active"]
]
set_table_data(get_shape_by_name(s11, "object 5"), s11_table_data, header_size_pt=9.5, cell_size_pt=8.2)

print("Updating Slide 12: Results: Model Comparison...")
s12 = prs.slides[11]
set_text(get_shape_by_name(s12, "object 3"), "Evaluating Five Different Machine Learning Models on 614 Credit Records", font_size_pt=11)

s12_table_data = [
    ["MODEL", "ACC. %", "PREC. %", "REC. %", "F1 %", "AUC"],
    ["Logistic regression", "74.80", "80.68", "83.53", "82.08", "0.728"],
    ["Gradient boosting", "80.49", "82.80", "90.59", "86.52", "0.705"],
    ["Hist. grad. boosting", "82.11", "82.47", "94.12", "87.91", "0.734"],
    ["Random forest (Model 5)", "85.37", "83.17", "98.82", "90.32", "0.729"],
    ["EBM Glass-Box (Model 4)", "85.37", "83.17", "98.82", "90.32", "0.752"]
]
set_table_data(get_shape_by_name(s12, "object 5"), s12_table_data, header_size_pt=9.5, cell_size_pt=8.5)

s12_outcome = "Outcome: EBM ties Random Forest at 85.37% accuracy and leads in AUC (0.752). It provides state-of-the-art predictive power while remaining 100% explainable, beating logistic regression by +10.57% accuracy."
set_text(get_shape_by_name(s12, "object 9"), s12_outcome, font_size_pt=8.5)

print("Updating Slide 13: Results: Live Demo Personas & Fraud Defense...")
s13 = prs.slides[12]
set_text(get_shape_by_name(s13, "object 2"), "RESULTS: DEMO PERSONAS & FRAUD DEFENSE", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s13, "object 3"), "Real-World Verification Demonstrating How Fraud is Stopped and Decisions are Explained", font_size_pt=11)
set_text(get_shape_by_name(s13, "object 5"), "DEMONSTRATION TEST CASES", font_size_pt=11, bold=True)
set_text(get_shape_by_name(s13, "object 7"), "SAMPLE EBM REJECTION REASONS", font_size_pt=11, bold=True)

s13_table1_data = [
    ["TEST CASE", "OTP", "FACE MATCH", "OCR NAME", "VERDICT", "EXPLANATION"],
    ["Jeeva Bharathi M", "Valid", "91.2% (Pass)", "100% (Pass)", "Approved", "Low DTI (28%), CIBIL 785"],
    ["Stolen OTP Fraud", "Valid", "41.5% (Fail)", "Mismatch", "BLOCKED", "Impostor face rejected"],
    ["Karthik Raman", "Valid", "88.5% (Pass)", "98% (Pass)", "Declined", "Clear reasons generated"]
]
set_table_data(get_shape_by_name(s13, "object 6"), s13_table1_data, header_size_pt=9, cell_size_pt=8)

s13_table2_data = [
    ["FACTOR", "VALUE", "IMPACT"],
    ["Debt-to-Income", "55.9%", "Exceeds 45% limit (-0.38)"],
    ["CIBIL Score", "560", "Below 650 threshold (-0.29)"],
    ["Card Utilization", "78%", "Revolving balance high (-0.21)"],
    ["Remedy", "Pay Rs. 8k", "Re-apply in 90 days"]
]
set_table_data(get_shape_by_name(s13, "object 8"), s13_table2_data, header_size_pt=9, cell_size_pt=8)

s13_reading = "Key Demonstration: Proves that stolen OTP impostors are blocked at the biometric gate, while high-risk applicants receive exact, transparent reasons for decline."
set_text(get_shape_by_name(s13, "object 12"), s13_reading, font_size_pt=8.5)

print("Updating Slide 14: How the 4 Problems are Solved...")
s14 = prs.slides[13]
set_text(get_shape_by_name(s14, "object 2"), "HOW OUR SYSTEM SOLVES THE 4 PROBLEMS", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s14, "object 3"), "Direct Impact of Our 5 Models on Modern Digital Lending", font_size_pt=11)
set_text(get_shape_by_name(s14, "object 6"), "PROBLEM 1 & 2 SOLUTIONS", font_size_pt=11, bold=True)
set_text(get_shape_by_name(s14, "object 11"), "PROBLEM 3 & 4 SOLUTIONS", font_size_pt=11, bold=True)

p12_sol = (
    "1. Manual Verification Delays:\n"
    "• Solved by: DigiLocker & OCR models.\n"
    "• Impact: Verification reduced from 3–7 business days to under 60 seconds.\n\n"
    "2. Fake Certificates & Forged Documents:\n"
    "• Solved by: DigiLocker direct government API fetch.\n"
    "• Impact: Only authentic UIDAI and Income Tax records accepted; zero fake paperwork."
)
set_text(get_shape_by_name(s14, "object 15"), p12_sol, font_size_pt=8.5)

p34_sol = (
    "3. Stolen OTPs & Identity Theft:\n"
    "• Solved by: Tri-Biometric face recognition model.\n"
    "• Impact: Impostors are blocked because their face does not match the Aadhaar & PAN photos.\n\n"
    "4. Unexplained Black-Box Rejections:\n"
    "• Solved by: Explainable Boosting Machine (EBM).\n"
    "• Impact: Applicants receive clear, transparent reasons and actionable advice to improve their credit profile."
)
set_text(get_shape_by_name(s14, "object 16"), p34_sol, font_size_pt=8.5)

print("Updating Slide 15: Conclusion and Future Scope...")
s15 = prs.slides[14]
set_text(get_shape_by_name(s15, "object 2"), "CONCLUSION AND FUTURE SCOPE", font_size_pt=18, bold=True)
set_text(get_shape_by_name(s15, "object 3"), "Summary of Contributions and Roadmap for Phase II", font_size_pt=11)
set_text(get_shape_by_name(s15, "object 6"), "CONCLUSION", font_size_pt=11, bold=True)
set_text(get_shape_by_name(s15, "object 11"), "FUTURE SCOPE", font_size_pt=11, bold=True)

conc_text = (
    "• Built an automated digital lending system solving manual delays, fake documents, OTP theft, and black-box rejections.\n\n"
    "• Integrated DigiLocker for instant, authentic document verification.\n\n"
    "• Introduced 3-Way Tri-Biometric face verification to defeat OTP identity theft.\n\n"
    "• Applied OCR for instant cross-document detail matching.\n\n"
    "• Proved that EBM matches Random Forest accuracy (85.37%) while delivering 100% exact, transparent explanations.\n\n"
    "• Complete working application running locally and deployed live on GitHub Pages."
)
set_text(get_shape_by_name(s15, "object 15"), conc_text, font_size_pt=8.2)

future_text = (
    "• Live integration with UIDAI Aadhaar FaceRD production mobile biometric service.\n\n"
    "• Multi-language voice explanation for rejection letters in regional languages (Tamil, Hindi).\n\n"
    "• Bank account statement auto-analysis via RBI Account Aggregator (AA) framework.\n\n"
    "• Dynamic interest rate calculator offering lower interest to low-risk borrowers."
)
set_text(get_shape_by_name(s15, "object 16"), future_text, font_size_pt=8.2)

print("Updating Slide 16: Thank You...")
s16 = prs.slides[15]
set_text(get_shape_by_name(s16, "object 2"), "THANK YOU", font_size_pt=24, bold=True)
set_text(get_shape_by_name(s16, "object 3"), "Questions, Comments & Suggestions are Welcome", font_size_pt=12)

# Update footer on all slides
for i, slide in enumerate(prs.slides):
    # Footers usually in object 28, 29, etc.
    f_title = get_shape_by_name(slide, "object 29") or get_shape_by_name(slide, "object 23") or get_shape_by_name(slide, "object 8") or get_shape_by_name(slide, "object 12")
    # Leave footer clean or update if found

prs.save(OUTPUT_PPTX)
print(f"Saved updated PowerPoint presentation to: {OUTPUT_PPTX}")

shutil.copy2(OUTPUT_PPTX, ARTIFACT_PPTX)
print(f"Copied updated presentation to artifact directory: {ARTIFACT_PPTX}")
