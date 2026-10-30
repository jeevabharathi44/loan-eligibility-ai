# -*- coding: utf-8 -*-
import pptx

prs = pptx.Presentation(r"C:\Users\Jeeva\.gemini\antigravity\brain\e905d853-ffc3-4e6a-aec2-a3c083d192f0\.user_uploaded\media_1793376831728_cd86e238.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"=== SLIDE {idx+1} ===")
    for sh in slide.shapes:
        if sh.has_text_frame:
            tf = sh.text_frame
            for p in tf.paragraphs[:2]:
                if p.text.strip():
                    fnames = [r.font.name for r in p.runs if r.font.name]
                    fsizes = [f"{r.font.size.pt}pt" for r in p.runs if r.font.size]
                    txt = p.text.strip().replace('\n', ' ')[:35].encode('ascii', errors='replace').decode('ascii')
                    print(f"  {sh.name}: '{txt}' | font={fnames} | size={fsizes}")
        elif sh.has_table:
            t = sh.table
            print(f"  {sh.name} [TABLE]: {len(t.rows)}x{len(t.columns)}")
            for r_idx in range(min(2, len(t.rows))):
                row = t.rows[r_idx]
                cell_info = []
                for c in row.cells:
                    p = c.text_frame.paragraphs[0] if c.text_frame.paragraphs else None
                    if p and p.runs:
                        sz = p.runs[0].font.size.pt if p.runs[0].font.size else "def"
                        fn = p.runs[0].font.name or "def"
                        cell_txt = c.text.strip()[:12].encode('ascii', errors='replace').decode('ascii')
                        cell_info.append(f"{cell_txt} ({fn} {sz}pt)")
                print(f"    row {r_idx}: " + " | ".join(cell_info))
